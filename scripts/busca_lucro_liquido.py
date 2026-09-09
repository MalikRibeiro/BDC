import re
import openpyxl
import pandas as pd
from pathlib import Path

def main():
    print("="*80)
    print("BUSCA PROFUNDA: LUCRO LÍQUIDO EM TODAS AS ABAS (PADRAO 2 E 4)")
    print("="*80)
    
    base_dir = Path(r"c:\Users\C807951\Desktop\BDC")
    processadas_dir = base_dir / "ENTRADAS" / "fichas" / "comercializadoras" / "processadas"
    
    # Carrega base Silver para pegar amostras
    silver_path = base_dir / "SAIDAS" / "silver" / "fichas_comercializadoras_extraidas" / "fichas_comercializadoras_extraidas.parquet"
    if not silver_path.exists():
        print("Base Silver não encontrada.")
        return
        
    df = pd.read_parquet(silver_path)
    
    regex_lucro = re.compile(r'lucro\s*l[ií]quido|resultado\s*l[ií]quido|lucro.*preju[ií]zo.*exerc[ií]cio', re.IGNORECASE)
    
    for versao in ["padrao_2", "padrao_4"]:
        amostras = df[df["versao_ficha"] == versao]["arquivo_nome"].dropna().unique()
        # Pega as 3 primeiras amostras físicas
        amostras_fisicas = [a for a in amostras if (processadas_dir / a).exists()][:3]
        
        print(f"\n[{versao}] Buscando em {len(amostras_fisicas)} amostras:")
        for arquivo in amostras_fisicas:
            caminho_arquivo = processadas_dir / arquivo
            print(f"\n  -> Arquivo: {arquivo}")
            try:
                wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
                encontrou = False
                for aba_nome in wb.sheetnames:
                    ws = wb[aba_nome]
                    for r in range(1, min(ws.max_row, 300) + 1):
                        for c in range(1, min(ws.max_column, 30) + 1):
                            cell = ws.cell(row=r, column=c)
                            if cell.value and isinstance(cell.value, str):
                                if regex_lucro.search(cell.value):
                                    encontrou = True
                                    # Pega o valor da vizinhança
                                    vizinhos = []
                                    if c < ws.max_column:
                                        v_dir = ws.cell(row=r, column=c+1).value
                                        if v_dir is not None: vizinhos.append(f"Dir:{v_dir}")
                                    if r < ws.max_row:
                                        v_baixo = ws.cell(row=r+1, column=c).value
                                        if v_baixo is not None: vizinhos.append(f"Abaixo:{v_baixo}")
                                        
                                    print(f"      [Encontrado na aba '{aba_nome}'] Célula {cell.coordinate}: '{cell.value}'")
                                    if vizinhos:
                                        print(f"         Vizinhos: {' | '.join(vizinhos)}")
                if not encontrou:
                    print("      NÃO ENCONTRADO em nenhuma aba!")
            except Exception as e:
                print(f"      Erro ao ler arquivo: {e}")

if __name__ == "__main__":
    main()
