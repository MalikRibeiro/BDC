import openpyxl
from pathlib import Path
import os

files_to_check = {
    "padrao_4": "ATVOS SANTA LUZIA 31072026.xlsx",
    "padrao_2": "SAFIRA COM_05_03_2021.xlsx",
    "padrao_6": "POLLARIX 03062026.xlsx",
    "padrao_7": "ENGIE 12052026.xlsx"
}

base_dir = Path(r"c:\Users\C807951\Desktop\BDC\ENTRADAS\fichas\comercializadoras")

for padrao, filename in files_to_check.items():
    paths = list(base_dir.rglob(filename))
    if not paths:
        print(f"File not found recursively: {filename}")
        continue
    
    path = paths[0]
    print(f"\n--- Checking {filename} ({padrao}) ---")
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        print(f"Failed to load: {e}")
        continue
    
    # Check TIPO_COMERCIALIZADORA in FichaIndividual
    if "FichaIndividual" in wb.sheetnames:
        sheet = wb["FichaIndividual"]
        print("TIPO_COMERCIALIZADORA candidates:")
        for row in range(1, 30):
            for col in range(1, 10):
                val = str(sheet.cell(row=row, column=col).value).strip().upper()
                if val and ("PURA" in val or "GRUPO" in val or "COMERCIALIZADORA" in val):
                    coord = sheet.cell(row=row, column=col).coordinate
                    print(f"  FichaIndividual!{coord}: {val}")
                    
        print("\nCONTROLADOR candidates:")
        for row in range(1, 30):
            for col in range(1, 15):
                val = str(sheet.cell(row=row, column=col).value).strip().upper()
                if val and ("POLLARIX" in val or "ENGIE" in val or "SAFIRA" in val or "ATVOS" in val):
                    coord = sheet.cell(row=row, column=col).coordinate
                    print(f"  FichaIndividual!{coord}: {val}")
    
    if padrao == "padrao_2" and "Para_Limite_Comercializadoras" in wb.sheetnames:
        sheet = wb["Para_Limite_Comercializadoras"]
        print("\nPara_Limite_Comercializadoras TIPO candidates:")
        for row in range(1, 10):
            for col in range(1, 15):
                val = str(sheet.cell(row=row, column=col).value).strip().upper()
                if val and ("PURA" in val or "GRUPO" in val or "COMERCIALIZADORA" in val):
                    coord = sheet.cell(row=row, column=col).coordinate
                    print(f"  Para_Limite_Comercializadoras!{coord}: {val}")
