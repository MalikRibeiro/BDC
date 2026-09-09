import pandas as pd
from pathlib import Path
import openpyxl
import re

def main():
    silver_path = Path("SAIDAS/silver/fichas_comercializadoras_extraidas/fichas_comercializadoras_extraidas.parquet")
    if silver_path.exists():
        df = pd.read_parquet(silver_path)
    else:
        df = pd.read_csv(str(silver_path).replace(".parquet", ".csv"), sep=";")
        
    pastas_raw = Path("ENTRADAS/fichas/comercializadoras/processadas")
    
    print("="*60)
    print("FRENTE 2: VERIFICAÇÃO DE DATA LIMITE (30/04/2025)")
    print("="*60)
    
    df["DATA_DF_DT"] = pd.to_datetime(df["DATA_DEMONSTRACAO_FINANCEIRA"], errors="coerce")
    limite = pd.Timestamp('2025-04-30')
    
    for versao in ["padrao_2", "padrao_4"]:
        df_v = df[df["versao_ficha"] == versao].dropna(subset=["DATA_DF_DT"])
        total = len(df_v)
        antes = (df_v["DATA_DF_DT"] < limite).sum()
        depois = (df_v["DATA_DF_DT"] >= limite).sum()
        print(f"[{versao}] Total de registros com data válida: {total}")
        print(f"  -> Antes de 30/04/2025 (Sem allow_semantic): {antes}")
        print(f"  -> Depois de 30/04/2025 (Com allow_semantic): {depois}\n")

    print("="*60)
    print("FRENTE 3: BUSCA PROFUNDA (GAP REAL E WD AGROINDUSTRIAL)")
    print("="*60)
    
    regex_a = re.compile(r'\b(FCO|ROA|ROE)\b|NOTA\s+BOARD|BOARD|NOTA\s+BUREAU|SCORE\s+BUREAU|\bBUREAU\b', re.IGNORECASE)
    
    for versao in ["padrao_2", "padrao_3", "padrao_4", "padrao_5"]:
        amostras = df[df["versao_ficha"] == versao]["arquivo_nome"].dropna().unique()
        amostras_fisicas = [a for a in amostras if (pastas_raw / a).exists()][:2]
        
        for arquivo in amostras_fisicas:
            caminho_arquivo = pastas_raw / arquivo
            print(f"\nBuscando 5 indicadores (gap real) em [{versao}] {arquivo}...")
            try:
                wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
                # Foca nas abas de Memória de Cálculo / Conf. Puras DRE e DFC, ou todas
                abas_alvo = [a for a in wb.sheetnames if "memória" in a.lower() or "calculo" in a.lower() or "cálculo" in a.lower() or "dre" in a.lower() or "dfc" in a.lower()]
                if not abas_alvo: abas_alvo = wb.sheetnames
                    
                encontrou_algo = False
                for aba_nome in abas_alvo:
                    ws = wb[aba_nome]
                    for r in range(1, min(ws.max_row, 300) + 1):
                        for c in range(1, min(ws.max_column, 25) + 1):
                            cell = ws.cell(row=r, column=c)
                            if cell.value and isinstance(cell.value, str):
                                if regex_a.search(cell.value):
                                    encontrou_algo = True
                                    val = str(cell.value).strip().replace('\n', ' ')
                                    if len(val) > 40: val = val[:37] + "..."
                                    print(f"  ENCONTRADO [{aba_nome}] {cell.coordinate}: '{val}'")
                if not encontrou_algo:
                    print("  -> CONFIRMADO: Nenhum dos 5 indicadores foi encontrado nas abas matemáticas.")
                wb.close()
            except Exception as e:
                print(f"Erro: {e}")

    print("\nBuscando TIPO_COMERCIALIZADORA em TODAS AS ABAS do arquivo WD AGROINDUSTRIAL (padrao_7)...")
    wd_files = [f for f in df[df["versao_ficha"] == "padrao_7"]["arquivo_nome"].unique() if "WD" in str(f).upper()]
    if wd_files and (pastas_raw / wd_files[0]).exists():
        arquivo_wd = wd_files[0]
        caminho_wd = pastas_raw / arquivo_wd
        regex_b = re.compile(r'TIPO\s+DE\s+COMERCIALIZADORA|CPURA|CGRUPO|\bPURA\b|\bGRUPO\b', re.IGNORECASE)
        try:
            wb = openpyxl.load_workbook(caminho_wd, data_only=True)
            encontrou = False
            for aba_nome in wb.sheetnames:
                ws = wb[aba_nome]
                for r in range(1, min(ws.max_row, 300) + 1):
                    for c in range(1, min(ws.max_column, 30) + 1):
                        cell = ws.cell(row=r, column=c)
                        if cell.value and isinstance(cell.value, str):
                            if regex_b.search(cell.value):
                                if "grupo econ" in str(cell.value).lower(): continue
                                encontrou = True
                                val = str(cell.value).strip().replace('\n', ' ')
                                print(f"  ENCONTRADO [{aba_nome}] {cell.coordinate}: '{val}'")
            if not encontrou:
                print("  -> CONFIRMADO: Campo totalmente OMITIDO da ficha física.")
            wb.close()
        except Exception as e:
            print(f"Erro: {e}")

if __name__ == "__main__":
    main()
