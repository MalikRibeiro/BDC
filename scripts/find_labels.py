import openpyxl
from pathlib import Path

files_to_check = {
    "padrao_6": "POLLARIX 03062026.xlsx",
    "padrao_7": "ENGIE 12052026.xlsx",
    "padrao_4": "ATVOS SANTA LUZIA 31072026.xlsx",
    "padrao_2": "SAFIRA COM_05_03_2021.xlsx"
}

base_dir = Path(r"c:\Users\C807951\Desktop\BDC\ENTRADAS\fichas\comercializadoras")

for padrao, filename in files_to_check.items():
    paths = list(base_dir.rglob(filename))
    if not paths: continue
    
    print(f"\n--- {padrao}: {filename} ---")
    wb = openpyxl.load_workbook(paths[0], data_only=True)
    
    for sheet_name in ["FichaIndividual", "Para_Limite_Comercializadoras"]:
        if sheet_name not in wb.sheetnames: continue
        sheet = wb[sheet_name]
        
        for row in range(1, 100):
            for col in range(1, 20):
                val = str(sheet.cell(row=row, column=col).value).strip().upper()
                if not val or val == 'NONE': continue
                
                # Check for CONTROLADOR
                if "CONTROLADOR" in val:
                    neighbor = str(sheet.cell(row=row, column=col+1).value).strip()
                    neighbor_below = str(sheet.cell(row=row+1, column=col).value).strip()
                    print(f"[{sheet_name}] Found '{val}' at row {row}, col {col}")
                    print(f"  -> Right neighbor: {neighbor}")
                    print(f"  -> Below neighbor: {neighbor_below}")
                
                # Check for TIPO / PURA / GRUPO
                if "TIPO" in val and "COMER" in val:
                    neighbor = str(sheet.cell(row=row, column=col+1).value).strip()
                    print(f"[{sheet_name}] Found '{val}' at row {row}, col {col}")
                    print(f"  -> Right neighbor: {neighbor}")
                    
                if "PURA" in val or "GRUPO" in val:
                    print(f"[{sheet_name}] Cell contains '{val}' at row {row}, col {col}")
