import pandas as pd
from pathlib import Path
import openpyxl
import re

def buscar_amostras():
    silver_path = Path("SAIDAS/silver/fichas_comercializadoras_extraidas/fichas_comercializadoras_extraidas.parquet")
    if silver_path.exists():
        df = pd.read_parquet(silver_path)
    else:
        df = pd.read_csv(str(silver_path).replace(".parquet", ".csv"), sep=";")

    pastas_raw = Path("ENTRADAS/fichas/comercializadoras/processadas")
    
    regex_a = re.compile(r'\b(FCO|ROA|ROE)\b|LUCRO\s+L[IÍ]QUIDO|LUCRO/PREJU[IÍ]ZO|RESULTADO\s+L[IÍ]QUIDO|FLUXO\s+DE\s+CAIXA\s+OPERACIONAL|FLUXO\s+DE\s+CAIXA\s+DAS\s+ATIVIDADES|CAIXA\s+L[IÍ]QUIDO\s+GERADO|NOTA\s+BOARD|BOARD|NOTA\s+BUREAU|SCORE\s+BUREAU|\bBUREAU\b', re.IGNORECASE)
    regex_b = re.compile(r'TIPO\s+DE\s+COMERCIALIZADORA|CPURA|CGRUPO|\bPURA\b|\bGRUPO\b', re.IGNORECASE)
    
    versoes = {
        "padrao_2": regex_a,
        "padrao_3": regex_a,
        "padrao_4": regex_a,
        "padrao_5": regex_a,
        "padrao_6": regex_b,
        "padrao_7": regex_b
    }
    
    for versao, regex_comp in versoes.items():
        print("="*90)
        print(f"ANALISANDO VERSÃO: {versao}")
        print("="*90)
        
        amostras = df[df["versao_ficha"] == versao]["arquivo_nome"].dropna().unique()
        amostras_fisicas = [a for a in amostras if (pastas_raw / a).exists()][:2]
        
        if not amostras_fisicas:
            print("Nenhuma amostra física encontrada.")
            continue
            
        for arquivo in amostras_fisicas:
            caminho_arquivo = pastas_raw / arquivo
            print(f"\n>>> ARQUIVO: {arquivo}")
            try:
                wb = openpyxl.load_workbook(caminho_arquivo, data_only=True)
                for aba_nome in wb.sheetnames:
                    ws = wb[aba_nome]
                    max_row = min(ws.max_row, 150)
                    max_col = min(ws.max_column, 25)
                    
                    for r in range(1, max_row + 1):
                        for c in range(1, max_col + 1):
                            cell = ws.cell(row=r, column=c)
                            if cell.value and isinstance(cell.value, str):
                                if regex_comp.search(cell.value):
                                    if "grupo econ" in cell.value.lower(): continue
                                    
                                    coord = cell.coordinate
                                    val = str(cell.value).strip().replace('\n', ' ')
                                    if len(val) > 40: val = val[:37] + "..."
                                    
                                    # Vizinhos Dir, Dir2, Baixo, BaixoDir
                                    vizinhos = []
                                    if c + 1 <= max_col:
                                        v = ws.cell(row=r, column=c+1).value
                                        if v is not None: vizinhos.append(f"Dir:{str(v)[:20]}")
                                    if c + 2 <= max_col:
                                        v = ws.cell(row=r, column=c+2).value
                                        if v is not None: vizinhos.append(f"Dir2:{str(v)[:20]}")
                                    if r + 1 <= max_row:
                                        v = ws.cell(row=r+1, column=c).value
                                        if v is not None: vizinhos.append(f"Abaixo:{str(v)[:20]}")
                                    if r + 1 <= max_row and c + 1 <= max_col:
                                        v = ws.cell(row=r+1, column=c+1).value
                                        if v is not None: vizinhos.append(f"AbaixoDir:{str(v)[:20]}")
                                        
                                    viz_str = " | ".join(vizinhos).replace('\n', ' ')
                                    if not viz_str: viz_str = "Vazio"
                                    
                                    print(f"  [{aba_nome}] {coord}: '{val}' --> VIZINHOS: {viz_str}")
                                    
                wb.close()
            except Exception as e:
                print(f"Erro lendo arquivo: {e}")

if __name__ == "__main__":
    buscar_amostras()
