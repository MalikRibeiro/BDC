"""Operações de leitura e fechamento seguro de workbooks Excel."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils import get_column_letter


def abrir_pasta(path: str | Path) -> Any:
    """Abre um workbook Excel em modo somente leitura."""
    return load_workbook(
        filename=Path(path),
        data_only=True,
        read_only=True,
        keep_links=False,
    )

def fechar_pasta(workbook: Any) -> None:
    """Fecha um workbook ignorando falhas de liberação."""
    if workbook is None:
        return
    try:
        workbook.close()
    except Exception:
        return

def ler_celula(worksheet: Worksheet, cell_ref: str, return_meta: bool = False) -> Any:
    """Lê o valor de uma célula a partir de uma referência A1 (ex: 'A19')."""
    if not cell_ref or cell_ref == "0":
        return (None, None) if return_meta else None
    val = worksheet[cell_ref].value
    if return_meta:
        return val, {"aba": worksheet.title, "celula": cell_ref}
    return val
