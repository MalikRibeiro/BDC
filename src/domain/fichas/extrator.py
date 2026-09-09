from __future__ import annotations

import re
import math
import time
import logging
import warnings
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple
from openpyxl.utils.datetime import from_excel
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple
from common.nulos import is_nulo_textual
from common.texto import normalizar_texto

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

logger = logging.getLogger(__name__)

class LeitorPlanilha:
    """Representação intermediária de uma planilha, para leitura independente de I/O."""
    def __init__(self, abas_grid: dict[str, list[tuple]], aba_ativa: str = None):
        self.abas_grid = abas_grid
        self.abas_nomes = list(abas_grid.keys())
        self.aba_ativa = aba_ativa if aba_ativa else (self.abas_nomes[0] if self.abas_nomes else "")

    @classmethod
    def do_workbook(cls, workbook) -> LeitorPlanilha:
        grid = {}
        aba_ativa = workbook.active.title if workbook.active else None
        for sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
            grid[sheet_name] = list(ws.iter_rows(min_row=1, max_row=150, min_col=1, max_col=30, values_only=True))
        return cls(grid, aba_ativa)

    def ler_celula(self, aba: str, cell_ref: str) -> Any:
        if aba not in self.abas_grid:
            return None
        try:
            row, col = coordinate_to_tuple(cell_ref)
            grid = self.abas_grid[aba]
            if 0 <= row - 1 < len(grid) and 0 <= col - 1 < len(grid[row - 1]):
                return grid[row - 1][col - 1]
            return None
        except Exception:
            return None

def analisar_data_com_seguranca(date_val: Any) -> Optional[datetime]:
    if isinstance(date_val, datetime):
        return date_val
    if isinstance(date_val, (int, float)):
        try:
            return from_excel(date_val)
        except Exception:
            return None
    if isinstance(date_val, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y%m%d", "%d-%m-%Y", "%Y/%m/%d"):
            try:
                return datetime.strptime(date_val.strip(), fmt)
            except ValueError:
                pass
    return None

def valor_extraido_limpo(val: Any, data_type: Optional[str] = None) -> Any:
    """Normalização puramente de extração."""
    if val is None:
        return None
    
    dt_str = str(data_type).lower() if data_type else ""
    
    if isinstance(val, str):
        s_upper = val.strip().upper()
        
        if not s_upper or s_upper.startswith("#") or is_nulo_textual(val):
            return None

    if any(t in dt_str for t in ("float", "num", "dec", "int", "moeda", "percent", "taxa", "valor", "score")):
        if isinstance(val, (int, float)):
            if math.isnan(val) or math.isinf(val):
                return None
            return float(val)
        
        clean = str(val).strip()
        is_percent = "%" in clean
        
        if clean.startswith("(") and clean.endswith(")"):
            clean = "-" + clean[1:-1].strip()
        
        clean = re.sub(r"[^\d\,\.\-eE+]", "", clean)
        
        if not clean:
            return None
            
        last_comma = clean.rfind(",")
        last_dot = clean.rfind(".")
        
        try:
            if last_comma > last_dot:
                clean = clean.replace(".", "").replace(",", ".")
            elif last_dot > last_comma:
                if "," in clean:
                    clean = clean.replace(",", "")
                else:
                    if clean.count(".") > 1:
                        clean = clean.replace(".", "")
                    else:
                        parts = clean.split(".")
                        if len(parts) > 1 and len(parts[1]) == 3 and not any(t in dt_str for t in ("percent", "taxa")) and not is_percent:
                            clean = clean.replace(".", "")
            
            val_float = float(clean)

            if is_percent:
                val_float = val_float / 100.0
                
            return val_float
        except ValueError:
            return None

    if any(t in dt_str for t in ("date", "data")):
        return analisar_data_com_seguranca(val)

    if isinstance(val, float) and val.is_integer():
        val = int(val) 
    return str(val).strip()

def _tipo_extraido_valido(val: Any, data_type: str, field_name: str = "") -> bool:
    if val is None:
        return False
        
    dt_str = str(data_type).lower() if data_type else ""
    fn_lower = field_name.lower()
    
    if any(t in dt_str for t in ("float", "num", "dec", "int", "moeda", "percent", "taxa", "valor", "score", "pd")):
        if not isinstance(val, (int, float)):
            return False
            
        if val > 10000000 and str(int(val)).startswith(("20", "31", "30", "01")):
            return False

        campos_nao_zeraveis = (
            "ativo_total", "passivo_circulante", "vendas_liquidas", 
            "probabilidade_default", "pd", "patrimonio_liquido", "rol"
        )
        if val == 0.0 and any(k in fn_lower for k in campos_nao_zeraveis):
            return False
            
        campos_indicadores = ("roa", "roe", "fco", "fco_rol", "margem")
        if val == -10.0 and any(k in fn_lower for k in campos_indicadores):
            return False

        if any(k in fn_lower for k in ("probabilidade", "pd")):
            if val < 0.0 or val > 1.0:
                return False

        campos_de_balanco = ("ativo", "passivo", "patrimonio", "receita", "rol", "lucro", "fluxo", "fco")
        if any(k in fn_lower for k in campos_de_balanco):
            if val > 50000000000.0 or val < -50000000000.0:
                return False
                
        return True

    if any(t in dt_str for t in ("date", "data")):
        if isinstance(val, datetime):
            if 1990 <= val.year <= datetime.now().year + 10:
                return True
        return False
        
    if isinstance(val, str):
        v = val.lower().strip()
        if not v or v in ("tipo", "valor", "data", "descrição", "ajustado", "-", "0"):
            return False
            
        if len(v) > 60 and "endereco" not in fn_lower and "endereço" not in fn_lower:
            return False
            
        if fn_lower in ("auditor", "empresa", "sigla") and v.replace(".", "").replace(",", "").isdigit():
            return False
            
        if "agencia" in fn_lower or "agência" in fn_lower:
            if not any(k in v for k in ("fitch", "mood", "s&p", "sp", "standard")): 
                return False
                
        if "nota" in fn_lower or "rating" in fn_lower:
            if len(v) > 5 or any(k in v for k in ("menor", "qualidade", "classificação", "agência", "risco")): 
                return False
                
        if "auditor" in fn_lower:
            if len(v) > 40: return False
            
    return True

def busca_omnidirecional(leitor: LeitorPlanilha, search_pattern: str, data_type: str, sheet_hint: str = None, field_name: str = "", offset_col: int = None, offset_row: int = None) -> Tuple[Any, dict]:
    try:
        regex = re.compile(search_pattern, re.IGNORECASE)
    except re.error:
        return None, {}

    sheet_names = leitor.abas_nomes
    if sheet_hint:
        hint_clean = str(sheet_hint).replace(" ", "").lower()
        sheet_names = sorted(sheet_names, key=lambda x: 0 if hint_clean in x.replace(" ", "").lower() else 1)

    for sheet_name in sheet_names:
        grid = leitor.abas_grid.get(sheet_name, [])
        for r_idx, row_tuple in enumerate(grid):
            for c_idx, cell_value in enumerate(row_tuple):
                if cell_value and isinstance(cell_value, str):
                    if regex.search(cell_value.strip()):
                        targets = []
                        if offset_col is not None or offset_row is not None:
                            o_col = int(offset_col) if offset_col is not None else 0
                            o_row = int(offset_row) if offset_row is not None else 0
                            targets.append((r_idx + o_row, c_idx + o_col))
                            
                        default_targets = [(r_idx, c_idx + offset) for offset in range(1, 7)]
                        default_targets.extend([(r_idx + offset, c_idx) for offset in range(1, 3)])
                        
                        for dt in default_targets:
                            if dt not in targets:
                                targets.append(dt)
                        
                        for tr, tc in targets:
                            if 0 <= tr < len(grid) and 0 <= tc < len(grid[tr]):
                                raw_val = grid[tr][tc]
                                cleaned_val = valor_extraido_limpo(raw_val, data_type)
                                
                                if cleaned_val is not None and _tipo_extraido_valido(cleaned_val, data_type, field_name):
                                    col_letter = get_column_letter(tc + 1)
                                    coord = f"{col_letter}{tr + 1}"
                                    return cleaned_val, {
                                        "celula_origem": coord,
                                        "aba_origem": sheet_name,
                                        "metodo": "omnidirectional_regex"
                                    }
    return None, {}

def extrair_registro(leitor: LeitorPlanilha, layout_schema: Dict[str, Any], master_catalog: Dict[str, Any] = None, allow_semantic: bool = True) -> Tuple[Dict[str, Any], List[dict[str, Any]]]:
    extracted_data = {}
    metadata_list = []
    
    fields_to_extract = {}
    max_score = 0.0
    gates_to_check = []
    
    if not master_catalog or "fields" not in master_catalog:
        raise ValueError("O master_catalog é obrigatório e deve conter 'fields'.")

    for mc_field, mc_config in master_catalog["fields"].items():
        if mc_config.get("nature") == "OBSERVED":
            fields_to_extract[mc_field] = dict(mc_config)
            max_score += float(mc_config.get("weight", 0))
            if mc_config.get("criticality") == "GATE_ENGINE":
                gates_to_check.append(mc_field)
        
    score_obtido = 0.0

    for field_name, field_config in fields_to_extract.items():
        data_type = field_config.get("data_type") or field_config.get("type")
        if not data_type:
            fn_lower = field_name.lower()
            if any(t in fn_lower for t in ("ativo", "passivo", "lucro", "patrimonio", "capital", "venda", "receita", "lair", "lajir", "fco", "fluxo", "probabilidade", "pd", "rol", "reserva", "imposto", "resultado", "score", "ac_pc", "at_pt", "mtm", "dividendos")):
                data_type = "float"
            elif any(t in fn_lower for t in ("data", "dt")):
                data_type = "date"
            else:
                data_type = "string"
                
        LEGADO_MAPEAMENTO = {
            "FCO": ["SCORE_FCO_ROL", "FCO_ROL", "MARGEM_FLUXO_CAIXA"],
            "ROA": ["SCORE_ROA"],
            "ROE": ["SCORE_ROE"],
            "PATRIMONIO_LIQUIDO": ["PL", "TOTAL_PATRIMONIO_LIQUIDO"],
            "ATIVO_TOTAL_AJUSTADO": ["ATIVO_TOTAL", "TOTAL_ATIVOS"],
            "ATIVO_CIRCULANTE_AJUSTADO": ["ATIVO_CIRCULANTE"],
            "PASSIVO_CIRCULANTE_AJUSTADO": ["PASSIVO_CIRCULANTE"],
            "PASSIVO_NAO_CIRCULANTE_FINANCEIRO_AJUSTADO": ["PASSIVO_NAO_CIRCULANTE", "PASSIVO_N_CIRCULANTE"],
            "FLUXO_DE_CAIXA_DAS_ATIVIDADES_OPERACIONAIS": ["FCO", "FLUXO_DE_CAIXA_OPERACIONAL"],
            "LUCRO_LIQUIDO": ["RESULTADO_LIQUIDO", "LUCRO_PREJUIZO"],
            "ROL": ["RECEITA_OPERACIONAL_LIQUIDA", "RECEITA_LIQUIDA"]
        }

        layout_field_config = layout_schema.get("field_map", {}).get(field_name)
        
        if not layout_field_config and field_name in LEGADO_MAPEAMENTO:
            for alias in LEGADO_MAPEAMENTO[field_name]:
                layout_field_config = layout_schema.get("field_map", {}).get(alias)
                if layout_field_config:
                    break
                    
        layout_field_config = layout_field_config or {}
        
        sheet_hint = layout_field_config.get("sheet") or field_config.get("sheet")
        offset_col = layout_field_config.get("offset_col")
        offset_row = layout_field_config.get("offset_row")
        
        celula_estatica = layout_field_config.get("value_cell") or layout_field_config.get("cell") or field_config.get("value_cell") or field_config.get("cell")
        
        val = None
        meta = {
            "campo": field_name,
            "aba_origem": None,
            "celula_origem": None,
            "metodo": "falha_extracao",
            "valor": None
        }
        if celula_estatica and str(celula_estatica).strip() not in ("0", ""):
            try:
                aba_estatica = leitor.aba_ativa
                if sheet_hint:
                    hint_clean = str(sheet_hint).replace(" ", "").lower()
                    for aba_real in leitor.abas_nomes:
                        if hint_clean in aba_real.replace(" ", "").lower():
                            aba_estatica = aba_real
                            break

                raw_val = leitor.ler_celula(aba_estatica, celula_estatica)
                clean_val = valor_extraido_limpo(raw_val, data_type)
                if clean_val is not None and _tipo_extraido_valido(clean_val, data_type, field_name):
                    val = clean_val
                    meta.update({
                        "aba_origem": aba_estatica,
                        "celula_origem": celula_estatica,
                        "metodo": "estatico_layout",
                        "valor": val
                    })
            except Exception:
                pass
                
        field_allow_semantic = field_config.get("allow_semantic", True)

        if val is None and allow_semantic and field_allow_semantic:
            patterns = field_config.get("search_patterns")
            if patterns and isinstance(patterns, list) and len(patterns) > 0:
                search_pattern = "|".join(patterns)
            else:
                search_pattern = field_name.replace("_", r"\s*")
                    
            val_dinamico, meta_inf = busca_omnidirecional(
                leitor, 
                search_pattern, 
                data_type, 
                sheet_hint, 
                field_name, 
                offset_col=offset_col,
                offset_row=offset_row
            )
            if val_dinamico is not None:
                val = val_dinamico
                meta.update({
                    "aba_origem": meta_inf.get("aba_origem"),
                    "celula_origem": meta_inf.get("celula_origem"),
                    "metodo": meta_inf.get("metodo", "omnidirectional_regex"),
                    "valor": val
                })

        value_mapping = field_config.get("value_mapping")
        if val is not None and value_mapping:
            val_upper = str(val).strip().upper()
            val = value_mapping.get(val_upper, val)
            
        field_enum = field_config.get("enum")
        if val is not None and field_enum and isinstance(field_enum, list):
            if str(val).strip().upper() not in [str(e).strip().upper() for e in field_enum if e is not None]:
                val = None
                meta["metodo"] = "rejeitado_por_enum"

        extracted_data[field_name] = val
        if meta["metodo"] != "falha_extracao":
            meta["valor"] = val
            metadata_list.append(meta)
            
        if val is not None:
            score_obtido += float(field_config.get("weight", 0))

    score = (score_obtido / max_score) * 100 if max_score > 0 else 0
    extracted_data["INTEGRIDADE_EXTRAIDA_PERCENTUAL"] = round(score, 2)
    
    falha_gate = False
    for gate in gates_to_check:
        if extracted_data.get(gate) is None:
            falha_gate = True
            logger.debug("[GATE_ENGINE] Layout candidato descartado. Campo %s (GATE) ausente.", gate)
            break
            
    extracted_data["_FALHA_GATE_CRITICO"] = falha_gate
    
    if falha_gate:
        score = 0.0
        extracted_data["INTEGRIDADE_EXTRAIDA_PERCENTUAL"] = 0.0
        logger.debug("[INTEGRIDADE] Layout candidato recusado: Falha no GATE Crítico.")
    elif score >= 40.0:
        logger.debug("[INTEGRIDADE] Layout candidato atingiu %.2f%% de integridade.", score)
    else:
        logger.debug("[INTEGRIDADE] Layout candidato recusado: baixa integridade (%.2f%%).", score)

    return extracted_data, metadata_list

def avaliar_vencedor_por_grid(leitor: LeitorPlanilha, layouts: Dict[str, Any], master_catalog: Dict[str, Any] = None) -> Tuple[Dict[str, Any], List[dict[str, Any]], str]:
    

    if not master_catalog or "fields" not in master_catalog:
        raise ValueError("master_catalog é obrigatório.")

    expected_tabs_raw = set()
    for _, l_schema in layouts.items():
        if "expected_tabs" in l_schema:
            expected_tabs_raw.update(l_schema["expected_tabs"])
            
    if not expected_tabs_raw:
        expected_tabs_raw = {
            "V0", "Para_Limite_Comercializadoras", "Premissas", 
            "FichaIndividual", "Memória de Cálculo", "Conf. Puras_DRE", 
            "Dados Gerais e Qualitativos", "DRE", "Dem.Fin."
        }
    
    def norm_tab(t: str) -> str:
        s = normalizar_texto(t, caixa_alta=True, remover_acentuacao=True)
        return s.replace(" ", "") if s else ""
        
    expected_tabs_norm = {norm_tab(t) for t in expected_tabs_raw}
    workbook_tabs_norm = {norm_tab(t) for t in leitor.abas_nomes}
    
    if not expected_tabs_norm.intersection(workbook_tabs_norm):
        logger.warning(f"[VETO] Documento rejeitado. Nenhuma aba bate com as abas de layout: {leitor.abas_nomes}")
        return {}, [], "DOC_001_ESTRUTURA_INCOMPATIVEL"

    aba_ativa = leitor.aba_ativa
    linhas_lidas = len(leitor.abas_grid.get(aba_ativa, []))
    colunas_lidas = len(leitor.abas_grid.get(aba_ativa, [])[0]) if linhas_lidas else 0
    
    logger.info("Extração (Schema): Arquivo lido. %s colunas, %s linhas identificadas na aba '%s'.", colunas_lidas, linhas_lidas, aba_ativa)

    data_df = None
    search_regex = r"DATA\s*DA\s*DEMONSTRA[CÇ][AÃ]O|DATA\s*BASE|DATA\s*DA\s*DF"
    if master_catalog and "DATA_DEMONSTRACAO_FINANCEIRA" in master_catalog.get("fields", {}):
        patterns = master_catalog["fields"]["DATA_DEMONSTRACAO_FINANCEIRA"].get("search_patterns")
        if patterns:
            search_regex = "|".join(patterns)
            
    val_date, _ = busca_omnidirecional(leitor, search_regex, "date", None, "DATA_DEMONSTRACAO_FINANCEIRA")
    
    if isinstance(val_date, datetime):
        data_df = val_date
        
    layouts_to_test = list(layouts.items())
    allow_semantic = True

    melhor_score = -1.0
    vencedor_dados = {}
    vencedor_meta = []
    vencedor_nome = "NENHUM"

    inicio_extracao = time.perf_counter()
    
    for layout_name, layout_schema in layouts_to_test:
        extracted, metadata = extrair_registro(leitor, layout_schema, master_catalog, allow_semantic=allow_semantic)
        score = extracted.get("INTEGRIDADE_EXTRAIDA_PERCENTUAL", 0)
        
        if score > melhor_score:
            melhor_score = score
            vencedor_dados = extracted
            vencedor_meta = metadata
            vencedor_nome = layout_name

    duracao_forca_bruta_ms = (time.perf_counter() - inicio_extracao) * 1000

    logger.info(
        "[CHAMPION] Torneio Força Bruta finalizado em %.1f ms. Vencedor: '%s' com score de %.2f%%.", 
        duracao_forca_bruta_ms, vencedor_nome, melhor_score
    )

    return vencedor_dados, vencedor_meta, vencedor_nome


def extrair_registro_do_vencedor(workbook: Any, layouts: Dict[str, Any], master_catalog: Dict[str, Any] = None) -> Tuple[Dict[str, Any], List[dict[str, Any]], str]:
    """Wrapper legado para manter compatibilidade com consumidores antigos."""
    leitor = LeitorPlanilha.do_workbook(workbook)
    return avaliar_vencedor_por_grid(leitor, layouts, master_catalog)