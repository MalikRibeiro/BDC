# -*- coding: utf-8 -*-
"""Serviço de extração de dados dinâmico e omnidirecional das fichas Excel."""

from __future__ import annotations

import re
import math
import logging
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils.datetime import from_excel
from openpyxl.utils import get_column_letter

from common.excel import ler_celula

logger = logging.getLogger(__name__)

CUTOFF_DATE_LAYOUT_CHANGE = datetime(2025, 4, 30)

# ==============================================================================
# DICIONÁRIO DE INTELIGÊNCIA SEMÂNTICA UNIVERSAL (FALLBACK CONSUMIDORES)
# Mantido apenas para garantir a retrocompatibilidade com fichas de consumidores
# que ainda não foram migradas para o Master Catalog.
# ==============================================================================
MAPA_SEMANTICO_INTELIGENTE = {
    "CNPJ": r"^\s*CNPJ\b(?!.*(?:BBCE|CONTROLADOR))",
    "SIGLA": r"SIGLA",
    "FCO": r"FCO|MARGEM\s*DE\s*FLUXO\s*DE\s*CAIXA",
    "PATRIMONIO_LIQUIDO": r"PATRIM[OÔ]NIO\s*L[IÍ]QUIDO",
    "LUCRO_LIQUIDO": r"LUCRO\s*L[IÍ]QUIDO|RESULTADO\s*L[IÍ]QUIDO|LUCRO/PREJUIZO DO EXERCICIO",
    "PROBABILIDADE_DEFAULT": r"PROBABILIDADE\s*DEFAULT|PD\b|PD\s*=",
    "DATA_DEMONSTRACAO_FINANCEIRA": r"DATA\s*DA\s*DEMONSTRA[CÇ][AÃ]O|DATA\s*BASE|DATA\s*DA\s*DF",
    "RECEITA_LIQUIDA": r"RECEITA\s*L[IÍ]QUIDA|VENDAS\s*L[IÍ]QUIDAS|ROL",
    "VENDAS_LIQUIDAS": r"VENDAS\s*L[IÍ]QUIDAS|ROL|RECEITA\s*OPERACIONAL\s*L[IÍ]QUIDA",
    "ATIVO_TOTAL": r"ATIVO\s*TOTAL",
    "PASSIVO_CIRCULANTE": r"PASSIVO\s*CIRCULANTE",
    "ATIVO_CIRCULANTE": r"ATIVO\s*CIRCULANTE",
    "LUCRO_BRUTO": r"LUCRO\s*BRUTO",
    "LAJIR": r"LAJIR|RESULTADO\s*OPERACIONAL",
    "LAIR": r"LAIR|LUCRO\s*ANTES\s*DO\s*IMPOSTO",
    "ATIVO_CIRCULANTE_FINANCEIRO": r"ATIVO\s*CIRCULANTE\s*FINANCEIRO",
    "PASSIVO_CIRCULANTE_FINANCEIRO": r"PASSIVO\s*CIRCULANTE\s*FINANCEIRO",
    "PASSIVO_NAO_CIRCULANTE_FINANCEIRO": r"PASSIVO\s*N[AÃ]O\s*CIRCULANTE\s*FINANCEIRO",
    "EMPRESA": r"EMPRESA|RAZ[AÃ]O\s*SOCIAL",
    "TIPO_COMERCIALIZADORA": r"TIPO\s*DE\s*COMERCIALIZADORA",
    "DATA_ADESAO_CCEE": r"DATA\s*DE\s*ADES[AÃ]O",
    "CODIGO_CCEE": r"C[OÓ]DIGO\s*CCEE",
    "DATA_CALCULO": r"DATA\s*DA\s*FICHA",
    "SCORE_BUREAU": r"SCORE\s*BUREAU|SCORE\b",
    "QUANTIDADE_RESTRITIVOS": r"QUANTIDADE\s*DE\s*RESTRITIVOS",
    "CAPITAL_SOCIAL": r"CAPITAL\s*SOCIAL",
    "LUCROS_ACUMULADOS": r"LUCROS\s*ACUMULADOS",
    "RESERVA_DE_LUCROS": r"RESERVA\s*DE\s*LUCROS",
    "FLUXO_DE_CAIXA_DAS_ATIVIDADES_OPERACIONAIS": r"FLUXO\s*DE\s*CAIXA\s*OPERACIONAL|CAIXA\s*L[IÍ]QUIDO\s*GERADO",
    "AGENCIA": r"AG[EÊ]NCIA",
    "NOTA_CREDITO": r"NOTA\s*DE\s*CR[EÉ]DITO",
    "ROL": r"RECEITA\s*OPERACIONAL\s*L[IÍ]QUIDA|ROL",
    "AC_PC": r"AC\s*/\s*PC|ATIVO\s*CIRCULANTE\s*/\s*PASSIVO\s*CIRCULANTE",
    "AT_PT": r"AT\s*/\s*PT|ATIVO\s*TOTAL\s*/\s*PASSIVO\s*TOTAL",
    "MTM_TOTAL_PL": r"MTM\s*TOTAL\s*/\s*PL|MTM\s*/\s*PATRIM[OÔ]NIO",
    "DIVIDENDOS_JCP_LUCRO_LIQUIDO": r"\(?DIVIDENDOS\s*\+\s*JCP\)?\s*/\s*LUCRO\s*L[IÍ]QUIDO|DIVIDENDOS\s*E\s*JCP",
    "CAPITAL_CIRCULANTE_LIQUIDO": r"CAPITAL\s*CIRCULANTE\s*L[IÍ]QUIDO|CCL\b",
    "RESTRITIVOS": r"RESTRITIVOS|APONTAMENTOS\s*RESTRITIVOS",
    "CNAE": r"CNAE\b|C[OÓ]DIGO\s*DE\s*ATIVIDADE",
    "NATUREZA_JURIDICA": r"NATUREZA\s*JUR[IÍ]DICA",
    "ENDERECO": r"ENDERE[CÇ]O|LOGRADOURO",
    "ROA": r"ROA\b|RETORNO\s*SOBRE\s*ATIVO",
    "ROE": r"ROE\b|RETORNO\s*SOBRE\s*PATRIM[OÔ]NIO",
    "FCO_ROL": r"FCO\s*/\s*ROL|FLUXO\s*DE\s*CAIXA\s*/\s*RECEITA",
    "CNPJ_BBCE": r"CNPJ\s*BBCE",
    "RATING_COPEL": r"RATING\s*COPEL",
    "RATING_PUBLICO": r"RATING\s*P[UÚ]BLICO",
    "SCORE_QUANTITATIVO": r"SCORE\s*QUANTITATIVO",
    "SCORE_QUALITATIVO": r"SCORE\s*QUALITATIVO"
}

def analisar_data_com_seguranca(date_val: Any) -> Optional[datetime]:
    if isinstance(date_val, datetime):
        return date_val
    if isinstance(date_val, (int, float)):
        try:
            return from_excel(date_val)
        except Exception:
            return None
    if isinstance(date_val, str):
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y%m%d", "%d-%m-%Y", "%Y/%m/%d"):
            try:
                return datetime.strptime(date_val.strip(), fmt)
            except ValueError:
                pass
    return None

def valor_extraido_limpo(val: Any, data_type: Optional[str] = None) -> Any:
    """Higieniza o valor extraído e garante o casting correto."""
    if val is None:
        return None
    
    dt_str = str(data_type).lower() if data_type else ""
    
    if isinstance(val, str):
        s_upper = val.strip().upper()
        if not s_upper or s_upper.startswith("#") or s_upper in ("NAN", "NONE", "<NA>", "N/A", "NULL", "N/D", "-", "--"):
            return None

    if any(t in dt_str for t in ("float", "num", "dec", "int", "moeda", "percent", "taxa", "valor", "score")):
        if isinstance(val, (int, float)):
            if math.isnan(val) or math.isinf(val):
                return None
            return float(val)
        
        clean = str(val).strip()
        # Notação contábil negativa (1.500) -> -1.500
        if clean.startswith("(") and clean.endswith(")"):
            clean = "-" + clean[1:-1].strip()
        
        # Limpa tudo que não for dígito, vírgula, ponto ou sinal de menos (remove R$, $, %, letras)
        clean = re.sub(r"[^\d\,\.-]", "", clean)
        
        if not clean:
            return None
            
        # Resolução de pontuação (milhar vs decimal)
        last_comma = clean.rfind(",")
        last_dot = clean.rfind(".")
        
        try:
            if last_comma > last_dot:
                # Padrão Brasileiro: 1.500,50 -> 1500.50
                clean = clean.replace(".", "").replace(",", ".")
            elif last_dot > last_comma:
                # Padrão Americano: 1,500.50 -> 1500.50
                if "," in clean:
                    clean = clean.replace(",", "")
                else:
                    # Só tem ponto: "1.500" ou "1.5"
                    if clean.count(".") > 1:
                        # Vários pontos: "1.500.000" -> "1500000"
                        clean = clean.replace(".", "")
                    else:
                        # Exatamente um ponto. Se tiver 3 dígitos depois do ponto, no Brasil quase sempre é milhar se a origem for string suja de excel.
                        # Exceções: taxas ou percentuais (onde 1.500 pode ser 1.5%)
                        parts = clean.split(".")
                        if len(parts[1]) == 3 and not any(t in dt_str for t in ("percent", "taxa")):
                            clean = clean.replace(".", "")
            
            return float(clean)
        except ValueError:
            return None

    if any(t in dt_str for t in ("date", "data")):
        return analisar_data_com_seguranca(val)

    if isinstance(val, float) and val.is_integer():
        val = int(val) 
    return str(val).strip()

def _tipo_extraido_valido(val: Any, data_type: str, field_name: str = "") -> bool:
    if val is None:
        return False
    dt_str = str(data_type).lower() if data_type else ""
    if any(t in dt_str for t in ("float", "num", "dec", "int", "moeda", "percent", "taxa", "valor", "score", "pd")):
        return isinstance(val, (int, float))
    if any(t in dt_str for t in ("date", "data")):
        return isinstance(val, datetime)
        
    # Sanity checks for strings to avoid grabbing headers or explanatory text
    if isinstance(val, str):
        v = val.lower().strip()
        if not v or v in ("tipo", "valor", "data", "descrição", "ajustado"):
            return False
        
        fn_lower = field_name.lower()
        
        # Rejeitar strings maiores que 60 chars (rodapés, observações), exceto se for endereço
        if len(v) > 60 and "endereco" not in fn_lower and "endereço" not in fn_lower:
            return False
            
        # Rejeitar números disfarçados de string em campos puramente de texto
        if fn_lower in ("auditor", "empresa", "sigla") and v.replace(".", "").replace(",", "").isdigit():
            return False
        
        # Heurísticas específicas por campo para evitar falsos positivos
        if "agencia" in fn_lower or "agência" in fn_lower:
            if not any(k in v for k in ("fitch", "mood", "s&p", "sp", "standard")): 
                return False
        if "nota" in fn_lower or "rating" in fn_lower:
            if len(v) > 5 or any(k in v for k in ("menor", "qualidade", "classificação", "agência", "risco")): 
                return False
        if "auditor" in fn_lower:
            if len(v) > 40: return False
            
    return True

def busca_omnidirecional(workbook: openpyxl.workbook.workbook.Workbook, search_pattern: str, data_type: str, sheet_hint: str = None, field_name: str = "", grid_cache: Dict[str, List[Tuple]] = None) -> Tuple[Any, dict]:
    """
    Caçador Universal (Refatorado para Performance in-memory RAM GRID):
    Varre TODAS as abas do Excel atrás da Regex através de uma matriz em memória.
    """
    if grid_cache is None:
        grid_cache = {}

    try:
        regex = re.compile(search_pattern, re.IGNORECASE)
    except re.error:
        return None, {}

    sheet_names = workbook.sheetnames
    if sheet_hint:
        hint_clean = str(sheet_hint).replace(" ", "").lower()
        sheet_names = sorted(sheet_names, key=lambda x: 0 if hint_clean in x.replace(" ", "").lower() else 1)

    for sheet_name in sheet_names:
        if sheet_name not in grid_cache:
            ws = workbook[sheet_name]
            # Convert worksheet to in-memory grid
            grid_cache[sheet_name] = list(ws.iter_rows(min_row=1, max_row=150, min_col=1, max_col=30, values_only=True))
            
        grid = grid_cache[sheet_name]
        
        for r_idx, row_tuple in enumerate(grid):
            for c_idx, cell_value in enumerate(row_tuple):
                if cell_value and isinstance(cell_value, str):
                    if regex.search(cell_value.strip()):
                        # Alvos: até 6 colunas à direita, e até 2 linhas abaixo
                        targets = [(r_idx, c_idx + offset) for offset in range(1, 7)]
                        targets.extend([(r_idx + offset, c_idx) for offset in range(1, 3)])
                        
                        for tr, tc in targets:
                            if 0 <= tr < len(grid) and 0 <= tc < len(grid[tr]):
                                raw_val = grid[tr][tc]
                                cleaned_val = valor_extraido_limpo(raw_val, data_type)
                                
                                if cleaned_val is not None and _tipo_extraido_valido(cleaned_val, data_type, field_name):
                                    col_letter = get_column_letter(tc + 1)
                                    coord = f"{col_letter}{tr + 1}"
                                    return cleaned_val, {
                                        "celula_origem": coord,
                                        "aba_origem": sheet_name,
                                        "metodo": "omnidirectional_regex"
                                    }
    return None, {}

def extrair_registro(workbook: openpyxl.workbook.workbook.Workbook, layout_schema: Dict[str, Any], master_catalog: Dict[str, Any] = None, grid_cache: Dict[str, List[Tuple]] = None) -> Tuple[Dict[str, Any], List[dict[str, Any]]]:
    extracted_data = {}
    metadata_list = []
    
    fields_to_extract = {}
    max_score = 0.0
    gates_to_check = []
    
    # 3. PROTEÇÃO AO LEGADO (Fallback)
    if master_catalog and "fields" in master_catalog:
        for mc_field, mc_config in master_catalog["fields"].items():
            if mc_config.get("nature") == "OBSERVED":
                fields_to_extract[mc_field] = dict(mc_config)
                max_score += float(mc_config.get("weight", 0))
                if mc_config.get("criticality") == "GATE_ENGINE":
                    gates_to_check.append(mc_field)
    else:
        # Fallback Consumidores
        field_map = layout_schema.get("field_map", {})
        fields_to_extract = dict(field_map)
        for sm_field in MAPA_SEMANTICO_INTELIGENTE.keys():
            if sm_field not in fields_to_extract:
                fields_to_extract[sm_field] = {}
        max_score = len(fields_to_extract) # each field weight = 1
        
    score_obtido = 0.0

    for field_name, field_config in fields_to_extract.items():
        data_type = field_config.get("data_type") or field_config.get("type")
        if not data_type:
            # Inferência de tipagem semântica para impedir que a busca omnidirecional aceite lixo (strings) no lugar de números
            fn_lower = field_name.lower()
            if any(t in fn_lower for t in ("ativo", "passivo", "lucro", "patrimonio", "capital", "venda", "receita", "lair", "lajir", "fco", "fluxo", "probabilidade", "pd", "rol", "reserva", "imposto", "resultado", "score", "ac_pc", "at_pt", "mtm", "dividendos")):
                data_type = "float"
            elif any(t in fn_lower for t in ("data", "dt")):
                data_type = "date"
            else:
                data_type = "string"
                
        sheet_hint = field_config.get("sheet")
        
        # 2. INTEGRAÇÃO COM O MASTER CATALOG E SCORING PONDERADO
        if master_catalog and "fields" in master_catalog:
            patterns = field_config.get("search_patterns")
            if patterns and isinstance(patterns, list) and len(patterns) > 0:
                search_pattern = "|".join(patterns)
            else:
                search_pattern = field_name.replace("_", r"\s*")
        else:
            search_pattern = MAPA_SEMANTICO_INTELIGENTE.get(field_name) or field_config.get("search_pattern")
            if not search_pattern:
                search_pattern = field_name.replace("_", r"\s*")
                
        val, meta_inf = busca_omnidirecional(workbook, search_pattern, data_type, sheet_hint, field_name, grid_cache)
        
        meta = {
            "campo": field_name,
            "aba_origem": meta_inf.get("aba_origem"),
            "celula_origem": meta_inf.get("celula_origem"),
            "metodo": meta_inf.get("metodo", "falha_extracao"),
            "valor": val
        }
        
        # Se a busca dinâmica falhar miseravelmente, tenta a coordenada fixa cega como último recurso
        # Isso ocorre apenas se não houver Master Catalog ou se o legacy mantiver coords.
        if val is None and not (master_catalog and "fields" in master_catalog):
            static_cell = field_config.get("value_cell") or field_config.get("cell")
            if static_cell and str(static_cell).strip() not in ("0", ""):
                try:
                    # Usa a aba sugerida no JSON ou a ativa
                    ws_estatico = workbook.active
                    if sheet_hint:
                        hint_clean = str(sheet_hint).replace(" ", "").lower()
                        for aba_real in workbook.sheetnames:
                            if hint_clean in aba_real.replace(" ", "").lower():
                                ws_estatico = workbook[aba_real]
                                break

                    raw_val, static_meta = ler_celula(ws_estatico, static_cell, return_meta=True)
                    clean_val = valor_extraido_limpo(raw_val, data_type)
                    if clean_val is not None and _tipo_extraido_valido(clean_val, data_type, field_name):
                        val = clean_val
                        meta.update({
                            "aba_origem": ws_estatico.title,
                            "celula_origem": static_cell,
                            "metodo": "estatico_fixo_fallback",
                            "valor": val
                        })
                except Exception:
                    pass

        extracted_data[field_name] = val
        if meta["metodo"] != "falha_extracao":
            metadata_list.append(meta)
            
        if val is not None:
            if master_catalog and "fields" in master_catalog:
                score_obtido += float(field_config.get("weight", 0))
            else:
                score_obtido += 1

    # Calcula a Integridade da Ficha
    score = (score_obtido / max_score) * 100 if max_score > 0 else 0
    extracted_data["INTEGRIDADE_EXTRAIDA_PERCENTUAL"] = round(score, 2)
    
    # 4. GATES de Segurança
    falha_gate = False
    for gate in gates_to_check:
        if extracted_data.get(gate) is None:
            falha_gate = True
            logger.warning(f"[GATE_ENGINE] Falha Crítica! Campo {gate} (GATE) ausente.")
            break
            
    extracted_data["_FALHA_GATE_CRITICO"] = falha_gate
    
    if falha_gate:
        # Penaliza severamente (zera o score) se o gate crítico falhou
        score = 0.0
        extracted_data["INTEGRIDADE_EXTRAIDA_PERCENTUAL"] = 0.0
        logger.warning("[INTEGRIDADE] Ficha recusada: Falha no GATE Crítico.")
    elif score >= 40.0:
        logger.info(f"[INTEGRIDADE] Ficha aprovada com {score:.2f}% de integridade (Score: {score_obtido}/{max_score}).")
    else:
        logger.warning(f"[INTEGRIDADE] Ficha recusada: apenas {score:.2f}% de integridade (Score: {score_obtido}/{max_score}).")

    return extracted_data, metadata_list

def extrair_registro_do_vencedor(
    workbook: openpyxl.workbook.workbook.Workbook,
    layouts: Dict[str, Any],
    master_catalog: Dict[str, Any] = None,
) -> Tuple[Dict[str, Any], List[dict[str, Any]], str]:
    """
    Motor Competitivo (Tournament Extraction):
    Ignora classificação cega baseada em uma única célula. 
    Testa a ficha contra TODOS os layouts e elege como 'Campeão' aquele que 
    atingir a maior integridade (porcentagem de campos com match válido).
    """
    from silver.normalizadores import normalizar_string

    # Veto por tipo de documento baseado em abas esperadas (Crítico #3)
    expected_tabs_raw = {
        "V0", "Para_Limite_Comercializadoras", "Premissas", 
        "FichaIndividual", "Memória de Cálculo", "Conf. Puras_DRE", 
        "Dados Gerais e Qualitativos", "DRE", "Dem.Fin."
    }
    
    def norm_tab(t: str) -> str:
        s = normalizar_string(t, upper=True)
        return s.replace(" ", "") if s else ""
        
    expected_tabs_norm = {norm_tab(t) for t in expected_tabs_raw}
    workbook_tabs_norm = {norm_tab(t) for t in workbook.sheetnames}
    
    if not expected_tabs_norm.intersection(workbook_tabs_norm):
        logger.warning(f"[VETO] Documento rejeitado. Nenhuma aba bate com as abas de layout: {workbook.sheetnames}")
        return {}, [], "DOC_001_ESTRUTURA_INCOMPATIVEL"

    best_score = -1.0
    champion_data = {}
    champion_meta = []
    champion_name = "NENHUM"
    
    grid_cache = {}

    for layout_name, layout_schema in layouts.items():
        extracted, metadata = extrair_registro(workbook, layout_schema, master_catalog, grid_cache, allow_semantic=True)
        score = extracted.get("INTEGRIDADE_EXTRAIDA_PERCENTUAL", 0)
        
        logger.info(f"Challenger {layout_name} obteve score: {score:.2f}%")
        
        if score > best_score:
            best_score = score
            champion_data = extracted
            champion_meta = metadata
            champion_name = layout_name

    logger.info(f"[CHAMPION] Torneio finalizado. Vencedor: '{champion_name}' com score de {best_score:.2f}%.")
    return champion_data, champion_meta, champion_name