"""Fase 0.7 — Auditoria somente leitura (sem conclusões textuais fixas).

Regras deste script:
- Não salva nenhuma ficha (.save nunca é chamado).
- Não engole exceções: todo erro é impresso com arquivo/etapa.
- Toda seção imprime N examinado / N população no terminal e no .md.
- O status final só é "CONCLUIDA" se todas as seções tiverem N > 0.
"""
from __future__ import annotations

import json
import math
import re
import sys
import traceback
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
from openpyxl.formula.tokenizer import Token, Tokenizer

PROJECT_ROOT = Path(__file__).resolve().parent.parent
SRC_DIR = PROJECT_ROOT / "src"
if str(SRC_DIR) not in sys.path:
    sys.path.insert(0, str(SRC_DIR))

from common.excel import abrir_pasta, fechar_pasta  # noqa: E402
from common.numeros import to_float_br  # noqa: E402
from domain.credito.pd_base import calcular_pd_base  # noqa: E402
from domain.fichas.extrator import (  # noqa: E402
    LeitorPlanilha,
    _tipo_extraido_valido,
    extrair_registro,
    valor_extraido_limpo,
)

DOCS = PROJECT_ROOT / "docs"
AMOSTRAS = DOCS / "amostras"
COM_DIR = PROJECT_ROOT / "ENTRADAS" / "fichas" / "comercializadoras"
CONS_DIR = PROJECT_ROOT / "ENTRADAS" / "fichas" / "consumidores"
SILVER_CSV = PROJECT_ROOT / "SAIDAS" / "silver" / "fichas_comercializadoras_extraidas" / "fichas_comercializadoras_extraidas.csv"
FATO_PARQUET = PROJECT_ROOT / "SAIDAS" / "relational" / "facts" / "credito" / "fato_analise_credito.parquet"
LOG_REJ = PROJECT_ROOT / "LOGS" / "rejeitados" / "2026-10-02_rejeicoes.json"
LAYOUT_V7 = PROJECT_ROOT / "ENTRADAS" / "control" / "layouts" / "layout_ficha_comercializadora_v7.json"
LAYOUT_V6 = PROJECT_ROOT / "ENTRADAS" / "control" / "layouts" / "layout_ficha_comercializadora_v6.json"
MASTER_CAT = PROJECT_ROOT / "ENTRADAS" / "control" / "quality" / "master_catalog_comercializadoras.json"
ZSCORE_CFG = PROJECT_ROOT / "ENTRADAS" / "control" / "configs" / "pd_zscore_config.json"
UNIVERSO = PROJECT_ROOT / "ENTRADAS" / "control" / "universo_pd.csv"

# Referência da NT §4 informada pelo usuário na diretriz da Fase 0.7 (não lida de arquivo).
NT_REF_USUARIO = {"intercept": -4.03, "coef_x12": -3.70, "coef_x16": 11.66, "coef_x19": -7.86, "coef_x22": -11.33}

SECOES_N: dict[str, int] = {}
MD: list[str] = []
ERROS: list[str] = []


def out(linha: str = "") -> None:
    print(linha)
    MD.append(linha)


def erro(etapa: str, arquivo: str, exc: BaseException) -> None:
    msg = f"[ERRO] etapa={etapa} arquivo={arquivo} tipo={type(exc).__name__} msg={exc}"
    print(msg)
    traceback.print_exc(limit=1)
    ERROS.append(msg)


def tabela_md(df: pd.DataFrame) -> None:
    if df.empty:
        out("(tabela vazia)")
        return
    print(df.to_string(index=False))
    MD.append("| " + " | ".join(map(str, df.columns)) + " |")
    MD.append("|" + "---|" * len(df.columns))
    for row in df.itertuples(index=False):
        MD.append("| " + " | ".join("" if pd.isna(v) else str(v).replace("|", "/").replace("\n", " ") for v in row) + " |")


def segmento_de(tipo: Any) -> str:
    t = "" if tipo is None or (isinstance(tipo, float) and math.isnan(tipo)) else str(tipo).strip().upper()
    if "PURA" in t:
        return "CPURA"
    if "GRUPO" in t:
        return "CGRUPO"
    return "INDEFINIDO"


def classificar_celula(raw_formula: Any, raw_cache: Any) -> str:
    e_form = isinstance(raw_formula, str) and raw_formula.startswith("=")
    if not e_form and (raw_cache is None or str(raw_cache).strip() == ""):
        return "VAZIA"
    if e_form and raw_cache is None:
        return "FORMULA_SEM_CACHE"
    zero = raw_cache in (0, 0.0, "0")
    if e_form and zero and "IFERROR" in raw_formula.upper():
        return "FORMULA_IFERROR_ZERO"
    if e_form:
        return "FORMULA_COM_CACHE"
    if zero:
        return "CONSTANTE_ZERO"
    return "CONSTANTE"


def refs_da_formula(formula: str) -> list[str]:
    tok = Tokenizer(formula)
    return [t.value.replace("$", "") for t in tok.items if t.type == Token.OPERAND and t.subtype == Token.RANGE]


def rastrear(wb_f, wb_v, sheet: str, cell: str, raiz: str, prof: int, max_prof: int, visit: set, linhas: list) -> None:
    chave = (sheet, cell)
    if chave in visit:
        return
    visit.add(chave)
    if sheet not in wb_f.sheetnames:
        linhas.append({"raiz": raiz, "prof": prof, "aba": sheet, "celula": cell, "tipo": "ABA_INEXISTENTE", "formula": "", "cache_presente": "", "number_format": ""})
        return
    cf = wb_f[sheet][cell]
    cv = wb_v[sheet][cell]
    tipo = classificar_celula(cf.value, cv.value)
    formula = cf.value if isinstance(cf.value, str) and cf.value.startswith("=") else ""
    linhas.append({
        "raiz": raiz, "prof": prof, "aba": sheet, "celula": cell, "tipo": tipo,
        "formula": formula, "cache_presente": "S" if cv.value is not None else "N",
        "cache_tipo": type(cv.value).__name__, "number_format": cf.number_format,
    })
    if not formula or prof >= max_prof:
        return
    for ref in refs_da_formula(formula):
        aba_ref, cel_ref = sheet, ref
        if "!" in ref:
            aba_ref, cel_ref = ref.rsplit("!", 1)
            aba_ref = aba_ref.strip("'")
        if ":" in cel_ref:
            linhas.append({"raiz": raiz, "prof": prof + 1, "aba": aba_ref, "celula": cel_ref, "tipo": "INTERVALO", "formula": "", "cache_presente": "", "number_format": ""})
            continue
        if not re.fullmatch(r"[A-Z]{1,3}\d{1,7}", cel_ref):
            linhas.append({"raiz": raiz, "prof": prof + 1, "aba": aba_ref, "celula": cel_ref, "tipo": "REF_NAO_CELULA", "formula": "", "cache_presente": "", "number_format": ""})
            continue
        rastrear(wb_f, wb_v, aba_ref, cel_ref, raiz, prof + 1, max_prof, visit, linhas)


def aba_ficha(nomes: list[str]) -> str | None:
    for n in nomes:
        if n.replace(" ", "").lower() == "fichaindividual":
            return n
    return None


def main() -> int:
    AMOSTRAS.mkdir(parents=True, exist_ok=True)
    out(f"# FASE 0.7 — AUDITORIA SOMENTE LEITURA ({datetime.now():%Y-%m-%d %H:%M:%S})")
    out(f"universo_pd.csv existe: {UNIVERSO.exists()}")

    silver = pd.read_csv(SILVER_CSV, sep=";", dtype=str)
    vig = silver[silver["_STATUS_REGISTRO"] == "VIGENTE"].copy()
    sub = silver[silver["_STATUS_REGISTRO"] != "VIGENTE"].copy()
    vig["SEG"] = vig["TIPO_COMERCIALIZADORA"].map(segmento_de)
    out(f"Silver: total={len(silver)} vigentes={len(vig)} nao_vigentes={len(sub)}")

    # ------------------------------------------------------------------ S1
    out("\n## [S1] RATING_COPEL vigente por segmento x layout")
    vig["RATING_LIMPO"] = vig["RATING_COPEL"].fillna("").str.strip()
    com_rt = vig[vig["RATING_LIMPO"] != ""]
    SECOES_N["S1_rating"] = len(com_rt)
    out(f"N com rating = {len(com_rt)} de {len(vig)} vigentes")
    ct = pd.crosstab([com_rt["SEG"], com_rt["versao_ficha"]], com_rt["RATING_LIMPO"], margins=True).reset_index()
    tabela_md(ct)

    # ------------------------------------------------------------------ S2
    out("\n## [S2] Rejeitados de comercializadoras reconstruídos por arquivo")
    arqs_rej = sorted(p.name for p in (COM_DIR / "rejeitadas").glob("*.xlsx"))
    arqs_rej_cons = {p.name for p in (CONS_DIR / "rejeitadas").glob("*.xlsx")} if (CONS_DIR / "rejeitadas").exists() else set()
    arqs_proc = {p.name for p in (COM_DIR / "processadas").glob("*.xlsx")}
    log = json.loads(LOG_REJ.read_text(encoding="utf-8"))
    por_arq = defaultdict(list)
    for ev in log:
        por_arq[ev.get("arquivo")].append(ev)
    origem = Counter()
    for nome in por_arq:
        if nome in arqs_rej:
            origem["comercializadoras/rejeitadas"] += 1
        elif nome in arqs_rej_cons:
            origem["consumidores/rejeitadas"] += 1
        elif nome in arqs_proc:
            origem["comercializadoras/processadas"] += 1
        else:
            origem["nao_localizado_em_disco"] += 1
    out(f"Eventos no log: {len(log)} | arquivos distintos no log: {len(por_arq)} | eventos repetidos (arquivo com >1 evento): {sum(1 for v in por_arq.values() if len(v) > 1)}")
    out(f"Arquivos distintos do log por localização atual: {dict(origem)}")
    linhas = []
    for nome in arqs_rej:
        evs = por_arq.get(nome, [])
        ult = evs[-1] if evs else {}
        linhas.append({
            "arquivo": nome, "n_eventos_log": len(evs),
            "ultimo_status": ult.get("status", "SEM_EVENTO_NO_LOG"),
            "ultimo_erro": "; ".join(ult.get("erros", []))[:120],
        })
    df_rej = pd.DataFrame(linhas)
    SECOES_N["S2_rejeitados"] = len(df_rej)
    out(f"N arquivos em comercializadoras/rejeitadas = {len(df_rej)}")
    tabela_md(df_rej["ultimo_status"].value_counts().rename_axis("ultimo_status").reset_index(name="n"))
    df_rej.to_csv(AMOSTRAS / "rejeitados_comercializadoras.csv", sep=";", index=False, encoding="utf-8-sig")
    out("Detalhe por arquivo: docs/amostras/rejeitados_comercializadoras.csv")

    # ------------------------------------------------------------------ S3
    out("\n## [S3] Decomposição dos registros não vigentes (797 = vigentes + não vigentes)")
    arqs_vig = set(vig["arquivo_nome"].dropna())
    sub["arquivo_tem_vigente"] = sub["arquivo_nome"].isin(arqs_vig)
    SECOES_N["S3_nao_vigentes"] = len(sub)
    tabela_md(sub["arquivo_tem_vigente"].value_counts().rename_axis("arquivo_tem_registro_vigente").reset_index(name="n_registros"))
    out(f"Arquivos distintos entre não vigentes: {sub['arquivo_nome'].nunique()}")
    out("Chave SCD2 usada pela Silver (orquestrador_comercializadoras.py:415): CNPJ + DATA_DEMONSTRACAO_FINANCEIRA")
    cols_hash = [c for c in silver.columns if "HASH" in c.upper()]
    out(f"Colunas de hash na Silver: {cols_hash}")

    # S3b: substituição com DATA_CALCULO mais recente no registro substituído
    out("\n### [S3b] Registros substituídos x vigente da mesma chave (CNPJ, DATA_DF)")
    k = ["CNPJ", "DATA_DEMONSTRACAO_FINANCEIRA"]
    comp_cols = [c for c in ["DATA_CALCULO", "RATING_COPEL", "PROBABILIDADE_DEFAULT", "DATA_BUREAU", "arquivo_nome"] + cols_hash if c in silver.columns]
    m = sub[k + comp_cols].merge(vig[k + comp_cols], on=k, how="left", suffixes=("_sub", "_vig"))
    m["dc_sub"] = pd.to_datetime(m.get("DATA_CALCULO_sub"), errors="coerce")
    m["dc_vig"] = pd.to_datetime(m.get("DATA_CALCULO_vig"), errors="coerce")
    res = {
        "n_substituidos": len(m),
        "sem_vigente_na_chave": int(m["arquivo_nome_vig"].isna().sum()),
        "mesmo_arquivo": int((m["arquivo_nome_sub"] == m["arquivo_nome_vig"]).sum()),
        "arquivo_diferente": int(((m["arquivo_nome_sub"] != m["arquivo_nome_vig"]) & m["arquivo_nome_vig"].notna()).sum()),
        "data_calculo_sub_mais_recente_que_vig": int((m["dc_sub"] > m["dc_vig"]).sum()),
    }
    for c in ["RATING_COPEL", "PROBABILIDADE_DEFAULT"]:
        if f"{c}_sub" in m.columns:
            res[f"{c}_diferente"] = int((m[f"{c}_sub"].fillna("") != m[f"{c}_vig"].fillna("")).sum())
    SECOES_N["S3b_substituicao"] = len(m)
    tabela_md(pd.DataFrame([res]))
    m.loc[m["dc_sub"] > m["dc_vig"], k + ["arquivo_nome_sub", "arquivo_nome_vig", "DATA_CALCULO_sub", "DATA_CALCULO_vig"]].to_csv(
        AMOSTRAS / "substituicao_data_calculo_mais_recente_perdida.csv", sep=";", index=False, encoding="utf-8-sig")
    dup_neg_com = [n for n, evs in por_arq.items() if any(e.get("status") == "ERRO_DUPLICIDADE_NEGOCIO" for e in evs) and (n in arqs_rej or n in arqs_proc)]
    out(f"Eventos ERRO_DUPLICIDADE_NEGOCIO cujo arquivo está em pastas de comercializadoras: {len(dup_neg_com)}")

    # ------------------------------------------------------------------ S4
    out("\n## [S4] padrao_7 CPURA: origem dos números")
    p7 = vig[(vig["versao_ficha"] == "padrao_7") & (vig["SEG"] == "CPURA")]
    arqs_p7 = sorted(p7["arquivo_nome"].dropna())
    out(f"População padrao_7 CPURA vigente: {len(arqs_p7)}")
    layout_v7 = json.loads(LAYOUT_V7.read_text(encoding="utf-8"))
    master = json.loads(MASTER_CAT.read_text(encoding="utf-8"))

    amostra_trace = arqs_p7[:4]
    trace_rows: list[dict] = []
    estado_rows: list[dict] = []
    replay_rows: list[dict] = []
    alvo_estado = ["B26", "B27", "B28", "B29", "Q24", "Q25", "Q26", "B49"] + [f"N{i}" for i in range(67, 77)]
    campos_replay = {"VENDAS_LIQUIDAS": "Q24", "LUCRO_LIQUIDO": "Q25", "FLUXO_DE_CAIXA_DAS_ATIVIDADES_OPERACIONAIS": "Q26",
                     "ATIVO_TOTAL": "G16", "TIPO_COMERCIALIZADORA": "B18", "RATING_COPEL": "B49", "PROBABILIDADE_DEFAULT": "B26"}
    rotulos_qual = Counter()

    for nome in arqs_p7:
        path = COM_DIR / "processadas" / nome
        try:
            wb_f = openpyxl.load_workbook(path, data_only=False)
            wb_v = openpyxl.load_workbook(path, data_only=True)
        except Exception as exc:  # noqa: BLE001
            erro("S4_abrir", nome, exc)
            continue
        aba = aba_ficha(wb_f.sheetnames)
        if aba is None:
            out(f"[S4] {nome}: aba FichaIndividual não encontrada. Abas: {wb_f.sheetnames}")
            wb_f.close(); wb_v.close()
            continue
        for c in alvo_estado:
            try:
                estado_rows.append({"arquivo": nome, "celula": c, "estado": classificar_celula(wb_f[aba][c].value, wb_v[aba][c].value),
                                    "cache_tipo": type(wb_v[aba][c].value).__name__})
            except Exception as exc:  # noqa: BLE001
                erro("S4_estado", f"{nome}!{c}", exc)
        for r in range(55, 100):
            for col in ("A", "F", "L"):
                v = wb_v[aba][f"{col}{r}"].value
                if isinstance(v, str) and re.search(r"board|auditor|bureau|risk", v, re.I):
                    rotulos_qual[f"{col}{r}={v.strip()[:40]}"] += 1
        if nome in amostra_trace:
            for raiz in ["B26", "B27", "B28", "B29", "Q24", "Q25", "Q26", "B49"]:
                try:
                    rastrear(wb_f, wb_v, aba, raiz, f"{nome}!{raiz}", 0, 4, set(), trace_rows)
                except Exception as exc:  # noqa: BLE001
                    erro("S4_rastreio", f"{nome}!{raiz}", exc)
        wb_f.close(); wb_v.close()

        # Replay do caminho real do pipeline (abrir_pasta read_only + LeitorPlanilha + extrair_registro)
        wb_ro = None
        try:
            wb_ro = abrir_pasta(path)
            leitor = LeitorPlanilha.do_workbook(wb_ro)
            extraido, _meta = extrair_registro(leitor, layout_v7, master)
            aba_ro = aba_ficha(leitor.abas_nomes)
            linha_silver = vig[vig["arquivo_nome"] == nome].iloc[0]
            for campo, cel in campos_replay.items():
                bruto = leitor.ler_celula(aba_ro, cel) if aba_ro else None
                dt = (master["fields"].get(campo, {}) or {}).get("data_type") or (master["fields"].get(campo, {}) or {}).get("type")
                limpo = valor_extraido_limpo(bruto, dt or ("string" if campo in ("TIPO_COMERCIALIZADORA", "RATING_COPEL") else "float"))
                replay_rows.append({
                    "arquivo": nome, "campo": campo, "celula": cel,
                    "no_catalogo_observed": (master["fields"].get(campo, {}) or {}).get("nature") == "OBSERVED",
                    "data_type_catalogo": dt,
                    "bruto_tipo": type(bruto).__name__, "bruto_presente": bruto is not None,
                    "limpo_presente": limpo is not None,
                    "valido_tipo": _tipo_extraido_valido(limpo, dt or "", campo) if limpo is not None else False,
                    "replay_extrair_registro_presente": extraido.get(campo) is not None,
                    "silver_presente": str(linha_silver.get(campo, "")).strip() not in ("", "nan", "None"),
                })
        except Exception as exc:  # noqa: BLE001
            erro("S4_replay_pipeline", nome, exc)
        finally:
            fechar_pasta(wb_ro)

    df_estado = pd.DataFrame(estado_rows)
    SECOES_N["S4_estado_celulas"] = df_estado["arquivo"].nunique() if not df_estado.empty else 0
    out(f"\n### [S4a] Estado das células (N arquivos = {SECOES_N['S4_estado_celulas']} de {len(arqs_p7)})")
    if not df_estado.empty:
        tabela_md(pd.crosstab(df_estado["celula"], df_estado["estado"]).reset_index())
        zeros = df_estado[df_estado["celula"].isin(["B26", "B27", "B28", "B29"]) & df_estado["estado"].isin(["FORMULA_IFERROR_ZERO", "CONSTANTE_ZERO"])]
        out("Arquivos com zero em B26..B29 (estado e tipo do cache):")
        tabela_md(zeros)
        df_bz = df_estado[df_estado["celula"].isin(["B26", "B27", "B28", "B29"])]
        out("Tipo do valor em cache de B26..B29 (str indica retorno textual do IFERROR):")
        tabela_md(pd.crosstab(df_bz["celula"], df_bz["cache_tipo"]).reset_index())
        df_estado.to_csv(AMOSTRAS / "p7_estado_celulas.csv", sep=";", index=False, encoding="utf-8-sig")

    df_trace = pd.DataFrame(trace_rows)
    SECOES_N["S4_rastreio"] = len(df_trace)
    out(f"\n### [S4b] Rastreamento de precedentes (N arquivos = {len(amostra_trace)}; N nós = {len(df_trace)})")
    if not df_trace.empty:
        df_trace.to_csv(AMOSTRAS / "p7_rastreio_precedentes.csv", sep=";", index=False, encoding="utf-8-sig")
        prim = df_trace[df_trace["raiz"].str.startswith(amostra_trace[0])]
        out(f"Nós do primeiro arquivo ({amostra_trace[0]}); demais em docs/amostras/p7_rastreio_precedentes.csv")
        tabela_md(prim[["raiz", "prof", "aba", "celula", "tipo", "cache_presente", "number_format", "formula"]].assign(formula=lambda d: d["formula"].str[:90]))
        out("Abas alcançadas pelo rastreio (todas as amostras):")
        tabela_md(df_trace["aba"].value_counts().rename_axis("aba").reset_index(name="n_nos"))

    df_replay = pd.DataFrame(replay_rows)
    SECOES_N["S4_replay"] = df_replay["arquivo"].nunique() if not df_replay.empty else 0
    out(f"\n### [S4c] Replay do extrator real x Silver (N arquivos = {SECOES_N['S4_replay']} de {len(arqs_p7)})")
    if not df_replay.empty:
        agg = df_replay.groupby("campo")[["bruto_presente", "limpo_presente", "valido_tipo", "replay_extrair_registro_presente", "silver_presente"]].sum().reset_index()
        agg.insert(1, "no_catalogo_observed", df_replay.groupby("campo")["no_catalogo_observed"].first().values)
        agg.insert(2, "data_type_catalogo", df_replay.groupby("campo")["data_type_catalogo"].first().values)
        tabela_md(agg)
        df_replay.to_csv(AMOSTRAS / "p7_replay_extrator.csv", sep=";", index=False, encoding="utf-8-sig")

    SECOES_N["S4_rotulos_qualitativos"] = len(rotulos_qual)
    out(f"\n### [S4d] Rótulos com board/auditor/bureau/risk em A/F/L 55..99 (contagem de arquivos)")
    tabela_md(pd.DataFrame(rotulos_qual.most_common(), columns=["endereco_rotulo", "n_arquivos"]))

    # ------------------------------------------------------------------ S5
    out("\n## [S5] padrao_7 INDEFINIDO: leitura física de B18 e evidências")
    p7i = vig[(vig["versao_ficha"] == "padrao_7") & (vig["SEG"] == "INDEFINIDO")]
    linhas5 = []
    for _, r in p7i.iterrows():
        nome = r["arquivo_nome"]
        path = COM_DIR / "processadas" / nome
        try:
            wb_f = openpyxl.load_workbook(path, data_only=False)
            wb_v = openpyxl.load_workbook(path, data_only=True)
            aba = aba_ficha(wb_f.sheetnames)
            g = lambda c: wb_v[aba][c].value  # noqa: E731
            linhas5.append({
                "cnpj": r.get("CNPJ"), "arquivo": nome, "data_df_silver": r.get("DATA_DEMONSTRACAO_FINANCEIRA"),
                "A18_rotulo": g("A18"), "B18_estado": classificar_celula(wb_f[aba]["B18"].value, g("B18")),
                "B18_formula": wb_f[aba]["B18"].value if isinstance(wb_f[aba]["B18"].value, str) and str(wb_f[aba]["B18"].value).startswith("=") else "",
                "B18_cache": g("B18"), "H12_agencia": g("H12"), "H13_nota": g("H13"),
                "B20": g("B20"), "B21": g("B21"), "B49_rating": g("B49"),
                "abas": "|".join(wb_f.sheetnames)[:150],
            })
            wb_f.close(); wb_v.close()
        except Exception as exc:  # noqa: BLE001
            erro("S5", nome, exc)
    df5 = pd.DataFrame(linhas5)
    SECOES_N["S5_indefinido"] = len(df5)
    out(f"N lidos = {len(df5)} de {len(p7i)}")
    if not df5.empty:
        tabela_md(df5["B18_estado"].value_counts().rename_axis("B18_estado").reset_index(name="n"))
        tabela_md(df5.drop(columns=["abas"]))
        df5.to_csv(AMOSTRAS / "p7_indefinido_b18.csv", sep=";", index=False, encoding="utf-8-sig")

    # ------------------------------------------------------------------ S6
    out("\n## [S6] pd_base.py: config real x referência NT informada pelo usuário")
    cfg = json.loads(ZSCORE_CFG.read_text(encoding="utf-8"))
    tabela_md(pd.DataFrame([{"parametro": k, "config": cfg.get(k), "nt_ref_usuario": v, "mesmo_sinal": (cfg.get(k) is not None and (cfg[k] > 0) == (v > 0)), "igual": cfg.get(k) == v} for k, v in NT_REF_USUARIO.items()]))

    out("\n### [S6b] calcular_pd_base real com insumos da Silver (padrao_6 CPURA vigente)")
    p6 = vig[(vig["versao_ficha"] == "padrao_6") & (vig["SEG"] == "CPURA")]
    out(f"População padrao_6 CPURA vigente: {len(p6)} | Silver tem coluna PD_BASE: {'PD_BASE' in silver.columns}")

    class _LogPrint:
        def warning(self, msg, *a):
            print("[pd_base.warning] " + (msg % a))

    linhas6 = []
    for _, r in p6.iterrows():
        reg = {k2: (None if pd.isna(v) else v) for k2, v in r.to_dict().items()}
        reg.pop("PD_BASE", None)
        try:
            pd_func = calcular_pd_base(reg, "CPURA", cfg, _LogPrint())
        except Exception as exc:  # noqa: BLE001
            erro("S6_calcular_pd_base", r["arquivo_nome"], exc)
            continue
        pd_decl = to_float_br(r.get("PROBABILIDADE_DEFAULT"))
        f = lambda c: to_float_br(r.get(c))  # noqa: E731
        at, vl = f("ATIVO_TOTAL"), f("VENDAS_LIQUIDAS")
        pd_nt = None
        if at and vl:
            x12 = ((f("LUCROS_ACUMULADOS") or 0) + (f("RESERVA_DE_LUCROS") or 0)) / at
            x16 = ((f("PASSIVO_CIRCULANTE_FINANCEIRO") or 0) + (f("PASSIVO_NAO_CIRCULANTE_FINANCEIRO") or 0)) / at
            x19 = ((f("ATIVO_CIRCULANTE") or 0) - (f("PASSIVO_CIRCULANTE") or 0)) / at
            x22 = ((f("ATIVO_CIRCULANTE_FINANCEIRO") or 0) - (f("PASSIVO_CIRCULANTE_FINANCEIRO") or 0)) / vl
            z = cfg["intercept"] + cfg["coef_x12"] * x12 + cfg["coef_x16"] * x16 + cfg["coef_x19"] * x19 + cfg["coef_x22"] * x22
            pd_nt = 1 / (1 + math.exp(-max(-20, min(z, 20))))
        linhas6.append({
            "arquivo": r["arquivo_nome"], "pd_declarada": pd_decl, "pd_funcao_real": pd_func, "pd_formula_exp_menos_z": pd_nt,
            "dif_funcao": None if pd_func is None or pd_decl is None else abs(pd_func - pd_decl),
            "dif_exp_menos_z": None if pd_nt is None or pd_decl is None else abs(pd_nt - pd_decl),
        })
    df6 = pd.DataFrame(linhas6)
    SECOES_N["S6_pd_base"] = len(df6)
    if not df6.empty:
        for tol in (1e-6, 1e-4):
            out(f"tol={tol}: função real reproduz {int((df6['dif_funcao'] < tol).sum())} de {df6['dif_funcao'].notna().sum()} | 1/(1+exp(-z)) reproduz {int((df6['dif_exp_menos_z'] < tol).sum())} de {df6['dif_exp_menos_z'].notna().sum()}")
        out(f"função real devolveu None: {int(df6['pd_funcao_real'].isna().sum())} | PD declarada ausente: {int(df6['pd_declarada'].isna().sum())}")
        df6.to_csv(AMOSTRAS / "p6_reproducao_pd_base.csv", sep=";", index=False, encoding="utf-8-sig")
        tabela_md(df6.sort_values("dif_funcao", na_position="last").head(15))

    out("\n### [S6c] PD_BASE na fato_analise_credito x PROBABILIDADE_DEFAULT da Silver (CPURA vigente)")
    if FATO_PARQUET.exists():
        fato = pd.read_parquet(FATO_PARQUET)
        out(f"Colunas relevantes na fato: {[c for c in fato.columns if c in ('PD_BASE', 'PD_METODO', 'SEGMENTO_PD', 'STATUS_CALCULO_PD', '_STATUS_REGISTRO')]}")
        f2 = fato.copy()
        if "_STATUS_REGISTRO" in f2.columns:
            f2 = f2[f2["_STATUS_REGISTRO"] == "VIGENTE"]
        if "SEGMENTO_PD" in f2.columns:
            f2 = f2[f2["SEGMENTO_PD"] == "CPURA"]
        f2 = f2.assign(_CNPJ=f2["CNPJ"].astype(str).str.replace(r"\D", "", regex=True).str.zfill(14))
        sv = vig.assign(_CNPJ=vig["CNPJ"].astype(str).str.replace(r"\D", "", regex=True).str.zfill(14), PD_SILVER=vig["PROBABILIDADE_DEFAULT"].map(to_float_br))
        mj = f2.merge(sv[["_CNPJ", "PD_SILVER", "versao_ficha"]].drop_duplicates("_CNPJ"), on="_CNPJ", how="left")
        pdb = pd.to_numeric(mj.get("PD_BASE"), errors="coerce")
        cat = pd.Series("PD_BASE_NULO", index=mj.index)
        cat[pdb.notna() & mj["PD_SILVER"].notna() & ((pdb - mj["PD_SILVER"]).abs() < 1e-9)] = "IGUAL_A_SILVER (lida)"
        cat[pdb.notna() & mj["PD_SILVER"].notna() & ((pdb - mj["PD_SILVER"]).abs() >= 1e-9)] = "DIFERENTE_DA_SILVER (recalculada)"
        cat[pdb.notna() & mj["PD_SILVER"].isna()] = "SILVER_SEM_PD (recalculada)"
        mj["ORIGEM_PD_BASE"] = cat
        SECOES_N["S6_fato"] = len(mj)
        out(f"N registros CPURA vigentes na fato = {len(mj)}")
        tabela_md(mj["ORIGEM_PD_BASE"].value_counts().rename_axis("origem").reset_index(name="n"))
        mj[["CNPJ", "versao_ficha", "PD_SILVER", "PD_BASE", "ORIGEM_PD_BASE"]].to_csv(AMOSTRAS / "fato_origem_pd_base.csv", sep=";", index=False, encoding="utf-8-sig")
    else:
        out(f"fato_analise_credito.parquet não encontrado em {FATO_PARQUET}")
        SECOES_N["S6_fato"] = 0

    # ------------------------------------------------------------------ status
    out("\n## STATUS DAS SEÇÕES")
    tabela_md(pd.DataFrame([{"secao": s, "N": n} for s, n in SECOES_N.items()]))
    vazias = [s for s, n in SECOES_N.items() if n == 0]
    out(f"Erros registrados: {len(ERROS)}")
    out("STATUS: CONCLUIDA" if not vazias and not ERROS else f"STATUS: INCOMPLETA | seções vazias: {vazias} | erros: {len(ERROS)}")
    out("Pendente por ausência de universo_pd.csv: recorte por CNPJ, cruzamento de rejeitados e do bug com o universo." if not UNIVERSO.exists() else "")

    (DOCS / "investigacao_fase_07.md").write_text("\n".join(MD), encoding="utf-8")
    print(f"\nRelatório: {DOCS / 'investigacao_fase_07.md'}")
    return 0 if not vazias and not ERROS else 2


if __name__ == "__main__":
    sys.exit(main())
