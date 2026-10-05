"""Mapeia as fichas oficiais de referência em docs/fichas oficiais."""
from __future__ import annotations

import sys
from pathlib import Path
import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parent.parent
FICHAS_OFICIAIS = PROJECT_ROOT / "docs" / "fichas oficiais"


def auditar_arquivo(arq_path: Path):
    print(f"\n{'=' * 80}")
    print(f"📄 ARQUIVO: {arq_path.name} ({arq_path.stat().st_size / 1024:.1f} KB)")
    print(f"{'=' * 80}")

    wb = openpyxl.load_workbook(arq_path, read_only=True, data_only=True)
    print(f"Abas presentes ({len(wb.sheetnames)}): {wb.sheetnames}")

    # Aba principal
    aba_principal = None
    for name in ["FichaIndividual", "Ficha Individual", "Dados Gerais"]:
        if name in wb.sheetnames:
            aba_principal = name
            break
    if not aba_principal and wb.sheetnames:
        aba_principal = wb.sheetnames[0]

    ws = wb[aba_principal]
    print(f"Aba examinada: '{aba_principal}'")

    celulas_chave = [
        ("A9", "B9"), ("A11", "B11"), ("A13", "B13"), ("A18", "B18"),
        ("A20", "B20"), ("A21", "B21"), ("A26", "B26"), ("C26", "C26"),
        ("A27", "B27"), ("A28", "B28"), ("A29", "B29"),
        ("A32", "C32"), ("A49", "B49"),
        ("F14", "G14"), ("F16", "G16"), ("F20", "G20"), ("F24", "G24"), ("F25", "G25"), ("F26", "G26"),
        ("L24", "Q24"), ("L25", "Q25"), ("L26", "Q26"),
        ("F67", "H67"), ("L67", "N67"), ("F72", "H72"), ("L72", "N72"),
    ]

    print(f"\n{'Endereço':<12} | {'Label':<28} | {'Valor Extraído':<30}")
    print(f"{'-' * 12}-+-{'-' * 28}-+-{'-' * 30}")

    for lbl_addr, val_addr in celulas_chave:
        try:
            lbl_v = ws[lbl_addr].value
        except Exception:
            lbl_v = None
        try:
            val_v = ws[val_addr].value
        except Exception:
            val_v = None
        lbl_str = str(lbl_v).strip() if lbl_v is not None else ""
        val_str = str(val_v).strip() if val_v is not None else "VAZIO"
        print(f"{val_addr:<12} | {lbl_str[:28]:<28} | {val_str[:30]:<30}")

    wb.close()


def main():
    if not FICHAS_OFICIAIS.exists():
        print(f"[-] Diretório não encontrado: {FICHAS_OFICIAIS}")
        return 1

    arquivos = sorted(FICHAS_OFICIAIS.glob("*.xlsx"))
    print(f"[+] Total de fichas oficiais encontradas: {len(arquivos)}")

    for arq in arquivos:
        auditar_arquivo(arq)

    return 0


if __name__ == "__main__":
    main()
